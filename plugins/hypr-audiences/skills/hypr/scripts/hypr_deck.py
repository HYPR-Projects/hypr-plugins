#!/usr/bin/env python3
"""Gera o deck "Audience Discovery" de um plano de audiências HYPR (.pptx),
seguindo a estrutura e identidade dos decks oficiais HYPR (FY26).

Uso: python3 hypr_deck.py plano.json saida.pptx
Requer: python-pptx

Texto com *asteriscos* vira destaque em azul HYPR.
"""
import json, sys, os, datetime, re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "..", "assets")

# ---- Tokens (decks HYPR FY26) ----
# ---- Tokens oficiais (HYPR Design System "Own the Journey 2026") ----
DARK = RGBColor(0x1C, 0x26, 0x2F)       # midnight: capa, divisórias, destaque
DARK2 = RGBColor(0x24, 0x30, 0x3A)      # painel sobre midnight
LIGHT = RGBColor(0xFA, 0xFA, 0xFA)      # cloud: único fundo claro
CARD = RGBColor(0xE5, 0xEB, 0xF2)       # mist: cards, tags
CARD_BLUE = RGBColor(0xC5, 0xEA, 0xF6)  # cyan-soft: card destacado
LINE = RGBColor(0xE0, 0xE3, 0xE6)       # hairline rgba(28,38,47,.12) sobre cloud
LINE_STRONG = RGBColor(0xCE, 0xD3, 0xD7)
BLUE = RGBColor(0x33, 0x97, 0xB9)       # cyan: uma palavra por título
BLUE_DEEP = RGBColor(0x24, 0x6C, 0x84)
INK = RGBColor(0x1C, 0x26, 0x2F)
BODY = RGBColor(0x53, 0x68, 0x72)       # slate
MUTED = RGBColor(0x78, 0x90, 0x9C)      # steel
MUTED_DARK = RGBColor(0x78, 0x90, 0x9C)
WHITE = RGBColor(0xFC, 0xFE, 0xFE)      # offwhite: texto sobre midnight
SOFT = RGBColor(0xC9, 0xD2, 0xDB)

F_LIGHT, F_REG, F_MED, F_SEMI = "Urbanist Light", "Urbanist", "Urbanist Medium", "Urbanist SemiBold"
F_MONO = "IBM Plex Mono"
# raios do DS (20px/14px/8px em frame 1920) convertidos para o slide de 13,333in
R_LG, R_MD, R_SM = 20 / 1920 * 13.333, 14 / 1920 * 13.333, 8 / 1920 * 13.333
EIXO_BY_LAYER = {"core": "Comportamento", "conquista": "Comportamento", "afinidade": "Afinidade", "expansao": "Lifestyle", "expansão": "Lifestyle", "proximidade": "Proximidade"}

W, H = Inches(13.333), Inches(7.5)
MX = Inches(96 / 1920 * 13.333)   # margem horizontal do DS (96px em 1920)
MY = Inches(64 / 1080 * 7.5)      # margem vertical (64px em 1080)
TRACK_DISPLAY = -0.02             # tracking de display (em)
GROWTH, GROWTH_DEEP, GROWTH_SOFT = RGBColor(0x4C, 0xB0, 0x50), RGBColor(0x01, 0x83, 0x76), RGBColor(0xE4, 0xEB, 0xA4)
DECLINE, DECLINE_DEEP, DECLINE_SOFT = RGBColor(0xEA, 0x1E, 0x63), RGBColor(0xB4, 0x10, 0x48), RGBColor(0xFF, 0xBE, 0xBF)
YEAR = str(datetime.date.today().year)


def E(v):
    """Garante EMU inteiro (o PowerPoint rejeita coordenadas com casa decimal)."""
    return Emu(int(round(v)))


def fmt_vol(v, short=False):
    v = float(v or 0)
    if v >= 1e6:
        s = f"{v/1e6:.1f}".replace(".", ",")
        if s.endswith(",0"): s = s[:-2]
        return (s, "M") if short else s + " mi"
    if v >= 1e3:
        return (f"{v/1e3:.0f}", "k") if short else f"{v/1e3:.0f} mil"
    return (f"{v:.0f}", "") if short else f"{v:.0f}"


def norm(t):
    import unicodedata
    return "".join(c for c in unicodedata.normalize("NFD", str(t)) if unicodedata.category(c) != "Mn").lower().strip()


def fmt_int(v):
    return f"{int(v or 0):,}".replace(",", ".")


def split_title(name):
    """Realça a última palavra relevante em azul: 'Visitantes de Postos Ipiranga' -> 'Visitantes de Postos *Ipiranga*'."""
    name = re.sub(r"\s*\(base \d\)$", "", name)
    words = name.split()
    if len(words) <= 1: return f"*{name}*"
    return " ".join(words[:-1]) + f" *{words[-1]}*"


class Deck:
    def __init__(self, plan):
        self.p = plan
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.n = 0
        self.all_auds = [dict(a, layer=l["name"]) for l in plan["layers"] for a in l["audiences"]]

    # ---------- primitivas ----------
    def _run(self, para, text, size, color, font, align=None, spacing=None):
        parts = text.split("*")
        for i, part in enumerate(parts):
            if not part: continue
            r = para.add_run(); r.text = part
            r.font.size = Pt(size); r.font.name = font
            r.font.color.rgb = BLUE if i % 2 else color
            if spacing is not None:
                sp = spacing * size if abs(spacing) < 0.2 else spacing   # <0.2 = em; senão pt
                r._r.get_or_add_rPr().set("spc", str(int(sp * 100)))
        if align is not None: para.alignment = align

    def text(self, s, text, x, y, w, h, size, color, font=F_REG, align=None, lines=None, line_spacing=1.15, anchor=MSO_ANCHOR.TOP, spacing=None):
        tb = s.shapes.add_textbox(E(x), E(y), E(w), E(h))
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(lines if lines is not None else [text]):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.line_spacing = line_spacing
            self._run(para, line, size, color, font, align, spacing)
        return tb

    def mono(self, s, text, x, y, w, size=6.5, color=MUTED, align=None, h=Inches(0.22)):
        return self.text(s, text.upper(), x, y, w, h, size, color, F_MONO, align=align, spacing=0.6)

    def rect(self, s, x, y, w, h, fill, radius=None, line=None, lw=0.75):
        x, y, w, h = E(x), E(y), E(w), E(h)
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
        if radius:
            shp.adjustments[0] = min(0.5, Inches(radius) / min(w, h)) if radius > 0.5 else radius
        if fill is None: shp.fill.background()
        else: shp.fill.solid(); shp.fill.fore_color.rgb = fill
        if line is not None: shp.line.color.rgb = line; shp.line.width = Pt(lw)
        else: shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def hline(self, s, x, y, w, color=LINE, pt=0.75):
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x), E(y), E(x + w), E(y))
        ln.line.color.rgb = color; ln.line.width = Pt(pt)
        return ln

    def pill(self, s, text, x, y, fill=None, line=None, color=INK, size=6.5, font=F_MONO, dot=False, h=Inches(0.26), pad=Inches(0.16), w=None):
        label = ("●  " if dot else "") + text.upper() if font == F_MONO else (("●  " if dot else "") + text)
        est = Inches(0.062 if font == F_MONO else 0.075) * len(label) * (size / 7)
        w = w or est + pad * 2
        shp = self.rect(s, x, y, w, h, fill, radius=0.5, line=line)
        tf = shp.text_frame; tf.margin_left = tf.margin_right = pad; tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
        self._run(tf.paragraphs[0], label, size, color, font, PP_ALIGN.CENTER, 0.6 if font == F_MONO else None)
        return w

    def tag_box(self, s, text, x, y, w, h=Inches(0.42), dark=False):
        shp = self.rect(s, x, y, w, h, None, radius=R_MD, line=DARK_LINE if dark else LINE_STRONG)
        tf = shp.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
        tf.margin_top = tf.margin_bottom = 0
        self._run(tf.paragraphs[0], text.upper(), 7, WHITE if dark else INK, F_SEMI, PP_ALIGN.CENTER, 0.4)

    def figure_slot(self, s, x, y, w, h, caption):
        """Espaço reservado para foto (DS): moldura hairline raio 20, dois ticks de canto, uma legenda mono."""
        self.rect(s, x, y, w, h, None, radius=R_LG, line=LINE_STRONG)
        t = Inches(0.22); c = LINE_STRONG
        for (cx, cy, dx, dy) in [(x + Inches(0.3), y + Inches(0.3), 1, 1), (x + w - Inches(0.3), y + h - Inches(0.3), -1, -1)]:
            self.hline(s, cx if dx > 0 else cx - t, cy, t, c)
            ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(cx), E(cy if dy > 0 else cy - t), E(cx), E(cy + t if dy > 0 else cy))
            ln.line.color.rgb = c; ln.line.width = Pt(0.75)
        self.mono(s, caption, x + Inches(0.3), y + h - Inches(0.6), w - Inches(0.6), 6, MUTED)

    # ---------- chrome ----------
    def slide(self, dark, tag, footer_mid, footer_left=None):
        s = self.prs.slides.add_slide(self.blank); self.n += 1
        bg = s.background.fill; bg.solid(); bg.fore_color.rgb = DARK if dark else LIGHT
        fg = WHITE if dark else INK
        self.pill(s, tag, MX, MY, None, line=(MUTED_DARK if dark else LINE_STRONG), color=(SOFT if dark else INK), size=6, dot=True, h=Inches(0.24))
        logo = os.path.join(ASSETS, "hypr_white.png" if dark else "hypr_dark.png")
        s.shapes.add_picture(logo, E(W - MX - Inches(1.0)), E(MY + Inches(0.02)), width=E(Inches(1.0)))
        fl = footer_left or "▪ FY26 Oficial"
        fc = MUTED_DARK if dark else MUTED
        self.mono(s, fl, MX, H - MY - Inches(0.1), Inches(4), 5.5, fc)
        self.mono(s, footer_mid, Inches(4.5), H - MY - Inches(0.1), Inches(4.33), 5.5, fc, align=PP_ALIGN.CENTER)
        self.mono(s, self.p.get("footer_right") or f"{self.p.get('region', 'Nacional')} · {YEAR}", W - MX - Inches(3), H - MY - Inches(0.1), Inches(3), 5.5, fc, align=PP_ALIGN.RIGHT)
        return s

    # ---------- slides ----------
    def cover(self):
        p = self.p
        s = self.slide(True, p.get("cover_tag", "Audience Discovery · " + YEAR), "Capa", "▪ FY26 Oficial")
        self.mono(s, p.get("cover_kicker") or f"{p['brand']}  ·  Audience Discovery", MX, Inches(3.25), Inches(8), 7, BLUE)
        title = p.get("cover_title") or f"Audiências para *{p['brand']}*."
        self.text(s, title, MX, Inches(3.55), Inches(11.5), Inches(1.9), 54, WHITE, F_LIGHT, line_spacing=1.0, spacing=TRACK_DISPLAY)
        sub = p.get("subtitle") or p.get("campaign", "")
        self.text(s, sub, MX, Inches(5.5), Inches(9), Inches(0.5), 15, SOFT, F_REG)
        meta = p.get("cover_meta") or " · ".join(x for x in [p.get("campaign_type", "Plano de audiências"), p.get("vertical", ""), p.get("region", "Nacional")] if x)
        self.mono(s, meta, MX, Inches(6.05), Inches(8), 6.5, MUTED_DARK)

    def summary(self, sections):
        s = self.slide(False, "00 Sumário", "Sumário")
        self.mono(s, f"Índice   {len(sections):02d} seções", MX, Inches(3.15), Inches(4), 6.5, MUTED)
        self.text(s, "O que preparamos\npara *vocês*.", MX, Inches(3.45), Inches(5.5), Inches(1.6), 36, INK, F_LIGHT, line_spacing=1.05, spacing=TRACK_DISPLAY)
        x, y, w = Inches(6.6), Inches(2.1), Inches(6)
        rh = Inches(0.62)
        self.hline(s, x, y, w)
        for i, name in enumerate(sections):
            yy = y + i * rh
            self.text(s, f"{i+1:02d}", x, yy + Inches(0.14), Inches(0.7), Inches(0.4), 18, BLUE, F_LIGHT)
            self.text(s, name, x + Inches(0.8), yy + Inches(0.19), w - Inches(0.8), Inches(0.4), 13, INK, F_REG)
            self.hline(s, x, yy + rh, w)

    def divider(self, num, title, sub):
        s = self.slide(True, "Audience Discovery · " + YEAR, f"Seção {num:02d}", "▪ FY26 Oficial")
        self.text(s, f"{num:02d}", MX, Inches(2.0), Inches(5), Inches(1.9), 96, BLUE, F_LIGHT, line_spacing=0.9, spacing=TRACK_DISPLAY)
        self.text(s, f"{title}*.*", MX, Inches(3.95), Inches(10), Inches(1.1), 54, WHITE, F_LIGHT, spacing=TRACK_DISPLAY)
        self.text(s, sub, MX, Inches(5.15), Inches(8), Inches(0.9), 16, SOFT, F_REG, line_spacing=1.3)

    def context(self):
        p = self.p
        s = self.slide(False, "01 Contexto", "O desafio")
        self.text(s, p.get("context_title") or f"{p['brand']} e seus concorrentes, tudo no mesmo *mapa*.", MX, Inches(1.35), Inches(10.5), Inches(1.4), 32, INK, F_LIGHT, line_spacing=1.05)
        cols = p.get("context_columns") or [
            ("Onde e quando ativar?", p.get("where", "")), ("Para quem?", p.get("who", "")), ("Como provar?", p.get("how", ""))]
        cw = (W - 2 * MX - Inches(0.8)) / 3
        y = Inches(4.55)
        for i, (head, body) in enumerate(cols):
            x = MX + i * (cw + Inches(0.4))
            self.hline(s, x, y, cw)
            self.text(s, f"{i+1:02d}", x, y + Inches(0.25), Inches(1), Inches(0.5), 26, BLUE, F_LIGHT)
            self.text(s, head, x, y + Inches(0.85), cw, Inches(0.35), 13, INK, F_REG)
            self.text(s, body, x, y + Inches(1.2), cw - Inches(0.3), Inches(0.9), 9.5, BODY, F_REG, line_spacing=1.35)
        # premissas (chips) discretas
        prem = p.get("premises") or []
        if prem:
            x = MX
            for t in prem[:4]:
                x += self.pill(s, t, x, Inches(3.3), CARD, None, BODY, 7, F_MED, h=Inches(0.3)) + Inches(0.12)

    def layers_overview(self):
        p = self.p
        layers = p["layers"]
        s = self.slide(False, "02 Audiências · Camadas", "Audience plan")
        total = len(self.all_auds); vol = sum(a.get("volume") or 0 for a in self.all_auds)
        tw = vol or 1
        aff = round(sum((a.get("affinity") or 0) * (a.get("volume") or 0) for a in self.all_auds) / tw)
        self.text(s, p.get("layers_title") or f"*{total} audiências* a partir do briefing.", MX, Inches(1.3), Inches(10.5), Inches(0.8), 32, INK, F_LIGHT)
        self.text(s, p.get("layers_sub") or "Cada camada recebe um papel e uma mensagem diferente. O Brand Affinity ordena quem entra primeiro.", MX, Inches(2.1), Inches(9.5), Inches(0.7), 13, BODY, F_REG, line_spacing=1.35)
        # número grande à direita
        num, suf = fmt_vol(vol, short=True)
        self.text(s, "", W - MX - Inches(3.2), Inches(1.3), Inches(3.2), Inches(0.8), 40, INK, F_LIGHT, align=PP_ALIGN.RIGHT, lines=[f"{num}"])
        tb = s.shapes[-1]; r = tb.text_frame.paragraphs[0].add_run(); r.text = suf; r.font.size = Pt(18); r.font.name = F_LIGHT; r.font.color.rgb = INK
        self.mono(s, f"volume bruto · {aff}% affinity média", W - MX - Inches(3.2), Inches(2.15), Inches(3.2), 6, MUTED, align=PP_ALIGN.RIGHT)
        n = len(layers); gap = Inches(0.25)
        cw = (W - 2 * MX - gap * (n - 1)) / n
        y, ch = Inches(3.1), Inches(3.55)
        for i, l in enumerate(layers):
            x = MX + i * (cw + gap)
            hi = (i == 0)
            self.rect(s, x, y, cw, ch, CARD if hi else None, radius=R_LG, line=None if hi else LINE_STRONG)
            self.mono(s, f"{i+1:02d} · {l['name']}", x + Inches(0.3), y + Inches(0.3), cw, 6.5, BLUE if hi else MUTED)
            self.text(s, l.get("headline") or l["name"], x + Inches(0.3), y + Inches(0.6), cw - Inches(0.5), Inches(0.7), 15, INK, F_REG, line_spacing=1.15)
            self.text(s, l.get("description", ""), x + Inches(0.3), y + Inches(1.3), cw - Inches(0.55), Inches(1.2), 9.5, BODY, F_REG, line_spacing=1.35)
            lv = sum(a.get("volume") or 0 for a in l["audiences"]); num, suf = fmt_vol(lv, short=True)
            tb = self.text(s, num, x + Inches(0.3), y + ch - Inches(1.0), Inches(2.2), Inches(0.7), 34, INK, F_LIGHT)
            r = tb.text_frame.paragraphs[0].add_run(); r.text = suf; r.font.size = Pt(14); r.font.name = F_LIGHT; r.font.color.rgb = INK
            self.mono(s, f"{len(l['audiences'])} audiências\nestimadas", x + Inches(1.75) if lv < 1e7 else x + Inches(2.1), y + ch - Inches(0.88), Inches(1.5), 5.5, MUTED, h=Inches(0.4))

    def audience_card(self, a, idx, total):
        p = self.p
        s = self.slide(False, f"Principais audiências · {idx:02d}/{total:02d}", "Audiências propostas", "▪ Fontes: HYPR Audiences · estimativas")
        # painel esquerdo
        px, py, pw, ph = MX, Inches(1.3), Inches(4.9), Inches(5.35)
        img = a.get("image")
        if img and os.path.exists(img):
            s.shapes.add_picture(img, E(px), E(py), width=E(pw), height=E(ph))
        else:
            self.figure_slot(s, px, py, pw, ph, a.get("image_caption") or f"[ {norm(a.get('sub_category') or 'foto').replace('&', '·')} · pdv ]")
        # coluna direita (slots 02 a 08 do contrato Audience Discovery)
        x = Inches(6.1); w = W - MX - x
        eixo = a.get("eixo") or EIXO_BY_LAYER.get(norm(a.get("layer", "")), "Comportamento")
        self.mono(s, eixo, x, Inches(1.3), Inches(2.2), 6, MUTED)
        self.mono(s, a.get("cluster") or a.get("sub_category", ""), x + Inches(1.9), Inches(1.3), Inches(4.5), 6, BLUE)
        tl = len(re.sub(r"\s*\(base \d\)$", "", a["name"]))
        self.text(s, split_title(a["name"]), x, Inches(1.55) if tl <= 34 else Inches(1.6), w, Inches(0.7), 30 if tl <= 26 else (24 if tl <= 34 else 20), INK, F_LIGHT)
        self.hline(s, x, Inches(2.3), w)
        # slot 04: devices estimados + praça
        num, suf = fmt_vol(a.get("volume"), short=True)
        tb = self.text(s, num, x, Inches(2.45), Inches(2.6), Inches(0.9), 48, INK, F_LIGHT)
        r = tb.text_frame.paragraphs[0].add_run(); r.text = suf; r.font.size = Pt(20); r.font.name = F_LIGHT; r.font.color.rgb = INK
        nx = x + Inches(0.62) * len(num) + Inches(0.6)
        self.mono(s, "Devices\nestimados", nx, Inches(2.75), Inches(1.2), 5.5, MUTED, h=Inches(0.4))
        self.pill(s, a.get("praca") or p.get("region", "Nacional"), nx + Inches(1.15), Inches(2.72), CARD, None, INK, 7, F_MED, h=Inches(0.28))
        # slot 05: três contagens de endereços por place (se o plano trouxer); senão endereços + affinity + ranking
        places = a.get("places") or []
        if len(places) >= 3:
            stats = [(fmt_int(pl.get("n")), pl.get("l", "")) for pl in places[:3]]
        else:
            stats = [(fmt_int(a.get("addresses")), "Endereços\nmapeados"), (f"{a.get('affinity') or 0}%", "Brand\nAffinity"), (f"#{a.get('rank', idx)}", f"no ranking\nde {len(self.all_auds)}")]
        sw = Inches(2.1); sy = Inches(3.45)
        for i, (v, lab) in enumerate(stats):
            sx = x + i * sw
            if i: self.rect(s, sx - Inches(0.25), sy + Inches(0.05), Emu(9525), Inches(0.75), LINE_STRONG)
            self.text(s, v, sx, sy, sw, Inches(0.5), 24, BLUE, F_LIGHT)
            self.mono(s, lab, sx, sy + Inches(0.5), sw, 5.5, MUTED, h=Inches(0.4))
        # slot 06: hooks (duas tags com +)
        tx, ty = x, Inches(4.45)
        for t in (a.get("tags") or a.get("hooks") or [])[:2]:
            tx += self.pill(s, "+ " + t, tx, ty, CARD, None, INK, 7, F_MED, h=Inches(0.28)) + Inches(0.1)
        # slot 07: texto curto (2 linhas) + loc
        self.rect(s, x, Inches(4.95), Emu(19050), Inches(0.72), BLUE)
        quote = a.get("why") or a.get("description", "")
        if len(quote) > 150: quote = quote[:147].rsplit(" ", 1)[0] + "..."
        self.text(s, quote, x + Inches(0.25), Inches(4.93), w - Inches(0.3), Inches(0.8), 12.5, INK, F_LIGHT, line_spacing=1.3)
        self.text(s, a.get("profile") or a.get("meta") or f"ID {a.get('id','')}", x + Inches(0.25), Inches(5.72), w, Inches(0.25), 8.5, MUTED, F_REG)
        # slot 08: redes mapeadas (até 4 tiles tipográficos; "off" = inventário não confirmado)
        nets = a.get("networks") or ([a["brand"]] if a.get("brand") else [])
        nets = [n for n in nets if n][:4]
        if nets:
            self.mono(s, "Redes mapeadas", x, Inches(6.05), w, 5.5, MUTED)
            bw = (w - Inches(0.15) * 3) / 4
            for i, n in enumerate(nets):
                off = isinstance(n, dict) and n.get("off")
                label = n["name"] if isinstance(n, dict) else n
                shp = self.rect(s, x + i * (bw + Inches(0.15)), Inches(6.27), bw, Inches(0.4), None, radius=R_MD, line=LINE_STRONG)
                if off: shp.line.dash_style = 4
                tf = shp.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False; tf.margin_top = tf.margin_bottom = 0
                self._run(tf.paragraphs[0], label.upper(), 7, MUTED if off else INK, F_SEMI, PP_ALIGN.CENTER, 0.4)

    def ranking_table(self, auds, part=None):
        p = self.p
        s = self.slide(False, "02 Audiências · Ranking", "Audience ranking", "▪ Fontes: HYPR Audiences · estimativas")
        self.text(s, (p.get("ranking_title") or "Ranking completo por *Brand Affinity*") + (f"  ·  {part}" if part else ""), MX, Inches(1.3), Inches(10), Inches(0.7), 28, INK, F_LIGHT)
        cx, cy, cw = MX, Inches(2.2), W - 2 * MX
        rh = Inches(0.37); ch = Inches(0.5) + rh * len(auds) + Inches(0.2)
        self.rect(s, cx, cy, cw, ch, RGBColor(0xFF, 0xFF, 0xFF), radius=R_LG, line=LINE_STRONG)
        cols = [(Inches(0.3), Inches(0.5), "#"), (Inches(0.8), Inches(3.9), "Audiência"), (Inches(4.8), Inches(1.6), "Camada"), (Inches(6.4), Inches(1.5), "Volume"), (Inches(7.9), Inches(1.4), "Endereços"), (Inches(9.3), Inches(2.3), "Brand Affinity")]
        for ox, ow, lab in cols:
            self.text(s, lab, cx + ox, cy + Inches(0.17), ow, Inches(0.3), 9, MUTED, F_MED)
        y = cy + Inches(0.5)
        for i, a in enumerate(auds):
            self.hline(s, cx + Inches(0.3), y, cw - Inches(0.6))
            yy = y + Inches(0.08)
            self.text(s, f"{a.get('rank', i+1):02d}", cx + Inches(0.3), yy, Inches(0.5), Inches(0.25), 9, BLUE, F_MED)
            self.text(s, re.sub(r"\s*\(base \d\)$", "", a["name"]), cx + Inches(0.8), yy, Inches(3.9), Inches(0.25), 9.5, INK, F_REG)
            self.text(s, a.get("layer", ""), cx + Inches(4.8), yy, Inches(1.6), Inches(0.25), 9, BODY, F_REG)
            self.text(s, fmt_vol(a.get("volume")), cx + Inches(6.4), yy, Inches(1.5), Inches(0.25), 9.5, INK, F_REG)
            self.text(s, fmt_int(a.get("addresses")), cx + Inches(7.9), yy, Inches(1.4), Inches(0.25), 9.5, INK, F_REG)
            if a.get("affinity") is None:
                self.text(s, "base: planos anteriores", cx + Inches(9.3), yy, Inches(2.3), Inches(0.25), 8.5, MUTED, F_REG)
            else:
                aff = max(0, min(100, a.get("affinity") or 0)); bw = Inches(1.6)
                self.rect(s, cx + Inches(9.3), yy + Inches(0.07), bw, Inches(0.09), CARD, radius=0.5)
                self.rect(s, cx + Inches(9.3), yy + Inches(0.07), max(Emu(1), int(bw * aff / 100)), Inches(0.09), BLUE if aff >= 75 else MUTED, radius=0.5)
                self.text(s, f"{aff}%", cx + Inches(11.0), yy - Inches(0.01), Inches(0.6), Inches(0.25), 9.5, INK, F_MED)
            y += rh

    def platform(self):
        p = self.p
        sols = {x["solution"]: x for x in p.get("solutions", [])}
        s = self.slide(False, "03 Plataforma", "Core products")
        order_all = [("geoIQ", "Onde", "Places Graph"), ("adsIQ", "Com o quê", "Max Attention"), ("revIQ", "Quanto", "Groundflow"), ("askIQ", "Por quê", None)]
        # só os pilares que se aplicam ao caso (applies=False ou priority "Não se aplica" ficam fora)
        order = [o for o in order_all if o[0] in sols and sols[o[0]].get("applies", True) and sols[o[0]].get("priority") != "Não se aplica"] or order_all
        n = len(order)
        self.mono(s, "Uma jornada · quatro IQs" if n == 4 else f"Uma jornada · {n} pilares", MX, Inches(1.3), Inches(5), 6.5, MUTED)
        self.text(s, p.get("platform_title") or "Como fecha o *ciclo*.", MX, Inches(1.55), Inches(10), Inches(0.8), 32, INK, F_LIGHT)
        self.text(s, p.get("platform_note") or "Os pilares rodam sobre a mesma base e se informam: o lugar orienta onde ativar, o formato garante atenção e a pesquisa explica o que ficou.", MX, Inches(2.4), Inches(9.5), Inches(0.7), 12, BODY, F_REG, line_spacing=1.35)
        gap = Inches(0.25); cw = (W - 2 * MX - gap * (n - 1)) / n; y, ch = Inches(3.45), Inches(3.2)
        body_size = 9.5 if n == 4 else 10.5
        for i, (name, q, pw) in enumerate(order):
            x = MX + i * (cw + gap); sol = sols.get(name, {})
            ess = sol.get("priority") == "Essencial"
            self.rect(s, x, y, cw, ch, CARD if ess else None, radius=R_LG, line=None if ess else LINE_STRONG)
            tile = self.rect(s, x + Inches(0.3), y + Inches(0.3), Inches(0.42), Inches(0.42), INK, radius=R_MD); tf = tile.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; self._run(tf.paragraphs[0], f"{i+1:02d}", 8, WHITE, F_SEMI, PP_ALIGN.CENTER)
            self.text(s, name, x + Inches(0.85), y + Inches(0.3), cw, Inches(0.4), 17, INK, F_REG)
            self.mono(s, q, x + Inches(0.3), y + Inches(0.78), cw, 6, BLUE)
            if pw: self.mono(s, f"powered by {pw}", x + Inches(0.3), y + Inches(0.98), cw, 5.5, MUTED)
            how = sol.get("how", "")
            self.text(s, how, x + Inches(0.3), y + Inches(1.35), cw - Inches(0.5), Inches(1.3), body_size, BODY, F_REG, line_spacing=1.35)
            pr = sol.get("priority")
            if pr: self.pill(s, pr, x + Inches(0.3), y + ch - Inches(0.6), CARD_BLUE if ess else None, None if ess else LINE_STRONG, BLUE_DEEP if ess else BODY, 6, F_MONO, h=Inches(0.24))
        ds = sols.get("Demandshift")
        if ds and ds.get("priority") in ("Essencial", "Recomendado"):
            self.text(s, f"*Demandshift*  ·  {ds.get('how', '')}", MX, Inches(6.75), Inches(11.5), Inches(0.3), 9, BODY, F_REG)

    MEASUREMENT_MODES = {
        "vendas": ("revIQ · powered by Groundflow", "A campanha virou *venda na loja*? E quanto.",
                   "Comparamos o consumo nos pontos de venda das regiões expostas à campanha com regiões equivalentes não expostas, e medimos o incremento no SKU anunciado e na categoria.",
                   [("Linha de base", "Vendas do SKU e da categoria antes do flight, região a região."), ("Exposto vs. controle", "Regiões impactadas comparadas a regiões equivalentes sem campanha."), ("Incremento", "A diferença entre os grupos é o efeito atribuível à mídia.")]),
        "visitas": ("geoIQ · powered by Places Graph", "A campanha levou gente *até o ponto*? Quantas.",
                    "Comparamos a taxa de visita aos pontos do anunciante entre os devices expostos à campanha e um grupo de controle equivalente não exposto, e medimos as visitas incrementais geradas pela mídia.",
                    [("Linha de base", "Taxa de visita aos pontos do anunciante antes do flight, por audiência e praça."), ("Exposto vs. controle", "Devices impactados comparados a um grupo equivalente sem campanha."), ("Visitas incrementais", "A diferença entre os grupos é o tráfego atribuível à mídia, com custo por visita incremental.")]),
        "marca": ("askIQ · Inteligência de Pesquisa", "A campanha mudou a *percepção da marca*?",
                  "Perguntamos a quem foi efetivamente exposto e a um grupo de controle equivalente, e medimos a diferença em lembrança, consideração e intenção: o brand lift atribuível à campanha.",
                  [("Pesquisa de base", "Questionário ao grupo de controle: lembrança, consideração e atributos da marca."), ("Exposto vs. controle", "O mesmo questionário aplicado a quem foi impactado pela campanha."), ("Brand lift", "A diferença entre os grupos é o efeito da mídia em cada métrica.")]),
    }

    def measurement(self):
        p = self.p
        mode = p.get("measurement_mode") or "vendas"
        label, title, approach, steps_default = self.MEASUREMENT_MODES.get(mode, self.MEASUREMENT_MODES["vendas"])
        s = self.slide(False, "03 Plataforma · Mensuração", "Como provar")
        self.mono(s, p.get("measurement_label") or label, MX, Inches(1.3), Inches(5), 6.5, MUTED)
        self.text(s, p.get("measurement_title") or title, MX, Inches(1.55), Inches(11.5), Inches(0.8), 32, INK, F_LIGHT)
        self.text(s, p.get("measurement") or approach, MX, Inches(2.5), Inches(5.6), Inches(1.6), 13, BODY, F_REG, line_spacing=1.4)
        steps = p.get("measurement_steps") or steps_default
        cw = (W - 2 * MX - Inches(0.8)) / 3; y = Inches(4.55)
        for i, (head, body) in enumerate(steps):
            x = MX + i * (cw + Inches(0.4))
            self.hline(s, x, y, cw)
            self.text(s, f"{i+1:02d}", x, y + Inches(0.25), Inches(1), Inches(0.5), 26, BLUE, F_LIGHT)
            self.text(s, head, x, y + Inches(0.85), cw, Inches(0.35), 13, INK, F_REG)
            self.text(s, body, x, y + Inches(1.2), cw - Inches(0.3), Inches(0.9), 9.5, BODY, F_REG, line_spacing=1.35)
        if p.get("measurement_secondary"):
            self.text(s, p["measurement_secondary"], MX, Inches(6.75), Inches(11.5), Inches(0.3), 9, MUTED, F_REG)

    def closing(self):
        p = self.p
        s = self.slide(True, "Próximos passos", "Próximos passos", "▪ FY26 Oficial")
        self.text(s, p.get("closing_title") or "Do bolso à *rua*.", MX, Inches(1.6), Inches(11), Inches(1.2), 54, WHITE, F_LIGHT, spacing=TRACK_DISPLAY)
        self.text(s, p.get("closing_subtitle") or "O que precisamos para ativar e medir este plano.", MX, Inches(2.8), Inches(9), Inches(0.5), 15, SOFT, F_REG)
        y = Inches(3.9)
        for i, st in enumerate((p.get("next_steps") or [])[:4]):
            self.hline(s, MX, y, Inches(8.5), MUTED_DARK)
            self.text(s, f"{i+1:02d}", MX, y + Inches(0.15), Inches(0.7), Inches(0.3), 12, BLUE, F_LIGHT)
            self.text(s, st, MX + Inches(0.75), y + Inches(0.17), Inches(7.7), Inches(0.35), 12.5, WHITE, F_REG)
            y += Inches(0.6)
        self.mono(s, "hyprgrid.ai", MX, Inches(6.5), Inches(3), 6.5, MUTED_DARK)

    def _clean(self):
        for sl in self.prs.slides:
            for shp in sl.shapes:
                st = shp._element.find(qn("p:style"))
                if st is not None: shp._element.remove(st)
                for el in shp._element.iter(qn("a:off"), qn("a:ext")):
                    for k, v in el.attrib.items():
                        if "." in v: el.set(k, str(int(round(float(v)))))

    def build(self, out):
        p = self.p
        max_cards = int(p.get("max_audience_cards", 4))
        ranked = sorted(self.all_auds, key=lambda a: -(a.get("affinity") or 0))
        for i, a in enumerate(ranked): a.setdefault("rank", i + 1)
        cards = ranked[:max_cards]
        mlabel = {"vendas": "Mensuração · venda incremental", "visitas": "Mensuração · visitas incrementais", "marca": "Mensuração · brand lift"}.get(p.get("measurement_mode") or "vendas", "Mensuração · exposto vs. controle")
        sections = ["Contexto & desafio", "Audiências propostas", "Plataforma HYPR · como fecha o ciclo", mlabel, "Próximos passos"]
        self.cover(); self.summary(sections)
        self.divider(1, "Contexto", p.get("context_sub") or p.get("objective", "")); self.context()
        self.divider(2, "Audiências", p.get("audiences_sub") or f"{len(self.all_auds)} recortes para {p['brand']}, em {len(p['layers'])} camadas, ordenados por Brand Affinity.")
        self.layers_overview()
        for i, a in enumerate(cards): self.audience_card(a, i + 1, len(cards))
        per = 12
        chunks = [ranked[i:i + per] for i in range(0, len(ranked), per)]
        for ci, ch in enumerate(chunks): self.ranking_table(ch, f"{ci+1}/{len(chunks)}" if len(chunks) > 1 else None)
        n_ok = len([x for x in p.get("solutions", []) if x.get("solution") in ("geoIQ", "adsIQ", "revIQ", "askIQ") and x.get("applies", True) and x.get("priority") != "Não se aplica"]) or 4
        self.divider(3, "Plataforma", p.get("platform_sub") or ("Os quatro IQs em jogo e como cada um fecha o ciclo deste plano." if n_ok == 4 else f"Os {n_ok} pilares em jogo neste plano e como cada um fecha o ciclo."))
        self.platform(); self.measurement(); self.closing()
        self._clean(); self.prs.save(out); return out


DARK_LINE = RGBColor(0x3A, 0x48, 0x53)

if __name__ == "__main__":
    if len(sys.argv) < 3: print(__doc__); sys.exit(1)
    plan = json.load(open(sys.argv[1], encoding="utf-8"))
    print(Deck(plan).build(sys.argv[2]))
